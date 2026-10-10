"""Legacy monthly review workbook (plans/plan-employee-month-sheet.md §7).

Reproduces the sheet HR used to fill by hand so admins keep their habit:
employees as column pairs ("Ngày" | "Nội dung"), sections as row blocks —
Quên điểm danh / Đi trễ-Về sớm / Nghỉ việc riêng / Nghỉ phép / Thiếu điểm danh.
"Nội dung" carries the reason already recorded in the system (explanations,
checkout-miss explanation, leave description); HR can add more in Excel.

:func:`legacy_layout` is pure (cells + merges + styles as data, unit-tested
without openpyxl); :func:`render_xlsx` turns it into bytes with openpyxl
(shipped with Frappe).
"""

from __future__ import annotations

import io

MIN_BLOCK_ROWS = 3  # keep each list block readable even when (almost) empty

# Style keys → used by render_xlsx.
LABEL = "label"  # black fill, white bold, centered (left column + names)
HEAD = "head"  # sub-header "Ngày" / "Nội dung"
DATE = "date"  # grey "Ngày" cell
TEXT = "text"  # plain "Nội dung" cell

SECTIONS = (
    ("forgot", "Quên điểm danh", "list"),
    ("late_early", "Đi trễ/Về sớm", "list"),
    ("leave_unpaid", "Nghỉ việc riêng | Số ngày - ngày (tháng)", "line"),
    ("leave_paid", "Nghỉ phép có lương | Số ngày - ngày (tháng)", "line"),
    ("absent", "Thiếu điểm danh cả ngày | Số ngày - ngày (tháng)", "line"),
)


def _line_text(item: dict) -> str:
    if not item or not item.get("text"):
        return ""
    count = item.get("count") or 0
    n = int(count) if float(count).is_integer() else count
    return f"{item['text']} ({n} ngày)"


def legacy_layout(employees: list[dict]) -> dict:
    """``employees``: ``[{"employee_name", "items": review_items(...)}]``.

    Returns ``{"cells": [(row, col, value, style)], "merges": [(r1, c1, r2, c2)],
    "widths": {col: width}, "rows": last_row}`` — 1-based like Excel.
    """
    cells: list[tuple] = []
    merges: list[tuple] = []
    n = len(employees)
    # Header rows 1–2.
    cells.append((1, 1, "", LABEL))
    cells.append((2, 1, "", LABEL))
    for i, emp in enumerate(employees):
        c = 2 + 2 * i
        cells.append((1, c, emp.get("employee_name") or emp.get("name") or "", LABEL))
        cells.append((1, c + 1, "", LABEL))
        merges.append((1, c, 1, c + 1))
        cells.append((2, c, "Ngày", HEAD))
        cells.append((2, c + 1, "Nội dung", HEAD))

    row = 3
    for key, label, kind in SECTIONS:
        if kind == "list":
            height = max([MIN_BLOCK_ROWS] + [len(e["items"].get(key) or []) for e in employees])
        else:
            height = 1
        cells.append((row, 1, label, LABEL))
        if height > 1:
            merges.append((row, 1, row + height - 1, 1))
        for r in range(row + 1, row + height):
            cells.append((r, 1, "", LABEL))
        for i, emp in enumerate(employees):
            c = 2 + 2 * i
            if kind == "list":
                entries = emp["items"].get(key) or []
                for k in range(height):
                    e = entries[k] if k < len(entries) else {}
                    cells.append((row + k, c, e.get("text", ""), DATE))
                    cells.append((row + k, c + 1, e.get("reason", ""), TEXT))
            else:
                item = emp["items"].get(key) or {}
                cells.append((row, c, _line_text(item), DATE))
                cells.append((row, c + 1, item.get("reason", ""), TEXT))
        row += height

    widths = {1: 30}
    for i in range(n):
        widths[2 + 2 * i] = 30
        widths[3 + 2 * i] = 26
    return {"cells": cells, "merges": merges, "widths": widths, "rows": row - 1}


def render_xlsx(layout: dict, sheet_title: str = "Bang cong") -> bytes:
    """Write :func:`legacy_layout` to an .xlsx (openpyxl)."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    red = Side(style="thin", color="C00000")
    border = Border(left=red, right=red, top=red, bottom=red)
    black = PatternFill("solid", fgColor="000000")
    grey = PatternFill("solid", fgColor="D9D9D9")
    styles = {
        LABEL: (
            black,
            Font(bold=True, color="FFFFFF", size=12),
            Alignment("center", "center", wrap_text=True),
        ),
        HEAD: (black, Font(bold=True, color="FFFFFF", size=9), Alignment("center", "center")),
        DATE: (grey, Font(size=9), Alignment("left", "center")),
        TEXT: (None, Font(size=9), Alignment("left", "center", wrap_text=True)),
    }

    wb = Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]
    for r, c, value, style in layout["cells"]:
        cell = ws.cell(row=r, column=c, value=value)
        fill, font, align = styles[style]
        if fill is not None:
            cell.fill = fill
        cell.font = font
        cell.alignment = align
        cell.border = border
    for r1, c1, r2, c2 in layout["merges"]:
        ws.merge_cells(start_row=r1, start_column=c1, end_row=r2, end_column=c2)
    for col, width in layout["widths"].items():
        ws.column_dimensions[get_column_letter(col)].width = width
    ws.row_dimensions[1].height = 22
    ws.freeze_panes = "B3"
    # Print: landscape, all employees across one page width.
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
